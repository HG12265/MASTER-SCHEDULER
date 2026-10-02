import io
import re
from datetime import datetime
from typing import Any, Dict, List, Optional, Set, Tuple
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from app.repositories.academic_year_repository import academic_year_repo
from app.repositories.class_repository import class_repo
from app.repositories.faculty_repository import faculty_repo
from app.repositories.institution_settings_repository import institution_settings_repo
from app.repositories.resource_repository import resource_repo
from app.repositories.semester_type_repository import semester_type_repo
from app.repositories.subject_repository import subject_repo
from app.repositories.time_slot_repository import time_slot_repo
from app.repositories.timetable_repository import timetable_entry_repo, timetable_repo
from app.repositories.working_day_repository import working_day_repo
from app.services.report_service import report_service
from app.utils.exceptions import NotFoundException


class ExcelExportService:
    def __init__(self):
        self.font_inst = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
        self.font_dept = Font(name="Calibri", size=11, bold=True, color="334155")
        self.font_title = Font(name="Calibri", size=11, bold=True, color="0F172A")
        self.font_meta = Font(name="Calibri", size=9, italic=True, color="475569")
        self.font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        self.font_cell_bold = Font(name="Calibri", size=9, bold=True)
        self.font_cell = Font(name="Calibri", size=9)
        self.font_footer = Font(name="Calibri", size=8, italic=True, color="64748B")

        self.fill_header = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        self.fill_sub_header = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
        self.fill_day = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        self.fill_break = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

        self.thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1"),
        )
        self.align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        self.align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    def _sanitize_sheet_title(self, title: str, existing_names: Set[str]) -> str:
        # Invalid Excel sheet chars: \ / ? * [ ] :
        clean = re.sub(r"[\\/*?:\[\]]", "_", title.strip())
        clean = clean[:31] if len(clean) > 31 else clean
        if not clean:
            clean = "Sheet"
        base = clean
        counter = 1
        while clean.lower() in existing_names:
            suffix = f"_{counter}"
            clean = f"{base[: 31 - len(suffix)]}{suffix}"
            counter += 1
        existing_names.add(clean.lower())
        return clean

    def _apply_autofit(self, ws, min_width=12, max_width=32):
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if "\n" in val:
                    lines = val.split("\n")
                    max_len = max(max_len, max(len(l) for l in lines))
                else:
                    max_len = max(max_len, len(val))
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, min_width), max_width)

    def _write_header_block(self, ws, inst: Dict[str, Any], subtitle: str, term_text: str, max_col: int):
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
        ws.cell(row=1, column=1, value=inst.get("institutionName", "Master University").upper()).font = self.font_inst
        ws.cell(row=1, column=1).alignment = self.align_center

        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max_col)
        ws.cell(row=2, column=1, value=inst.get("departmentName", "Department of Computer Science")).font = self.font_dept
        ws.cell(row=2, column=1).alignment = self.align_center

        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=max_col)
        ws.cell(row=3, column=1, value=f"{inst.get('academicTitle', 'CLASS TIME TABLE')} — {subtitle}").font = self.font_title
        ws.cell(row=3, column=1).alignment = self.align_center

        ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=max_col)
        ws.cell(row=4, column=1, value=term_text).font = self.font_meta
        ws.cell(row=4, column=1).alignment = self.align_center

    async def generate_class_timetable_excel(self, timetable_id: str, class_id: str) -> io.BytesIO:
        tt = await timetable_repo.get_by_id(timetable_id)
        if not tt:
            raise NotFoundException(f"Timetable with ID '{timetable_id}' not found")
        cls_doc = await class_repo.get_by_id(class_id)
        if not cls_doc:
            raise NotFoundException(f"Class with ID '{class_id}' not found")

        inst = await institution_settings_repo.get_settings()
        working_days = await working_day_repo.find_many({"isActive": True, "isWorkingDay": True}, sort=[("dayOrder", 1)])
        time_slots = await time_slot_repo.find_many({"isActive": True}, sort=[("slotOrder", 1)])
        entries = await timetable_entry_repo.find_many({"timetableId": timetable_id, "classId": class_id, "isActive": True})

        subjects = await subject_repo.find_many({"isActive": True})
        subj_map = {s["id"]: s for s in subjects}
        faculty_all = await faculty_repo.find_many({"isActive": True})
        fac_map = {f["id"]: f for f in faculty_all}
        resources = await resource_repo.find_many({"isActive": True})
        res_map = {r["id"]: r for r in resources}

        grid_map = {(e.get("workingDayId"), e.get("timeSlotId")): e for e in entries}

        wb = Workbook()
        ws = wb.active
        ws.title = self._sanitize_sheet_title(cls_doc.get("name", "Class Timetable"), set())

        ay_doc = await academic_year_repo.get_by_id(tt["academicYearId"])
        st_doc = await semester_type_repo.get_by_id(tt["semesterTypeId"])
        ay_name = ay_doc.get("name", "") if ay_doc else ""
        st_name = st_doc.get("name", "") if st_doc else ""
        term_text = f"Academic Year: {ay_name} | {st_name} | Version: v{tt.get('version', 1)} ({tt.get('status', 'DRAFT')})"

        num_cols = len(time_slots) + 1
        self._write_header_block(ws, inst, cls_doc.get("displayName") or cls_doc.get("name", "Class"), term_text, num_cols)

        # Header row
        ws.row_dimensions[6].height = 28
        cell_c1 = ws.cell(row=6, column=1, value="DAY / PERIOD")
        cell_c1.font = self.font_header
        cell_c1.fill = self.fill_header
        cell_c1.alignment = self.align_center
        cell_c1.border = self.thin_border

        for col_idx, ts in enumerate(time_slots, start=2):
            s_name = ts.get("name", "")
            s_time = f"{ts.get('startTime', '')} - {ts.get('endTime', '')}"
            c = ws.cell(row=6, column=col_idx, value=f"{s_name}\n({s_time})")
            c.font = self.font_header
            c.fill = self.fill_header
            c.alignment = self.align_center
            c.border = self.thin_border

        # Body rows
        cur_row = 7
        for wd in working_days:
            ws.row_dimensions[cur_row].height = 36
            d_cell = ws.cell(row=cur_row, column=1, value=wd.get("name", "Day"))
            d_cell.font = self.font_cell_bold
            d_cell.fill = self.fill_day
            d_cell.alignment = self.align_center
            d_cell.border = self.thin_border

            for col_idx, ts in enumerate(time_slots, start=2):
                c = ws.cell(row=cur_row, column=col_idx)
                c.border = self.thin_border
                c.alignment = self.align_center

                if ts.get("slotType") in ("BREAK", "LUNCH"):
                    c.value = ts.get("name", "")
                    c.font = self.font_meta
                    c.fill = self.fill_break
                    continue

                entry = grid_map.get((wd["id"], ts["id"]))
                if entry:
                    sid = entry.get("subjectId")
                    s_code = subj_map.get(sid, {}).get("subjectCode", entry.get("title", "")) if sid else entry.get("title", "")
                    fac_codes = [fac_map[fid].get("facultyCode") or fac_map[fid].get("name", "")[:10] for fid in entry.get("facultyIds", []) if fid in fac_map]
                    fac_str = ", ".join(fac_codes)
                    rid = entry.get("resourceId")
                    room_str = f" [{res_map[rid].get('code')}]" if rid and rid in res_map else ""

                    c.value = f"{s_code}\n{fac_str}{room_str}" if fac_str else s_code
                    c.font = self.font_cell_bold
                else:
                    c.value = "—"
                    c.font = self.font_meta
            cur_row += 1

        self._apply_autofit(ws)
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    async def generate_faculty_timetable_excel(self, timetable_id: str, faculty_id: str) -> io.BytesIO:
        tt = await timetable_repo.get_by_id(timetable_id)
        if not tt:
            raise NotFoundException(f"Timetable with ID '{timetable_id}' not found")
        fac_doc = await faculty_repo.get_by_id(faculty_id)
        if not fac_doc:
            raise NotFoundException(f"Faculty with ID '{faculty_id}' not found")

        inst = await institution_settings_repo.get_settings()
        working_days = await working_day_repo.find_many({"isActive": True, "isWorkingDay": True}, sort=[("dayOrder", 1)])
        time_slots = await time_slot_repo.find_many({"isActive": True}, sort=[("slotOrder", 1)])
        entries = await timetable_entry_repo.find_many({"timetableId": timetable_id, "facultyIds": faculty_id, "isActive": True})

        subjects = await subject_repo.find_many({"isActive": True})
        subj_map = {s["id"]: s for s in subjects}
        classes = await class_repo.find_many({"isActive": True})
        cls_map = {c["id"]: c for c in classes}
        resources = await resource_repo.find_many({"isActive": True})
        res_map = {r["id"]: r for r in resources}

        grid_map = {(e.get("workingDayId"), e.get("timeSlotId")): e for e in entries}

        wb = Workbook()
        ws = wb.active
        ws.title = self._sanitize_sheet_title(fac_doc.get("name", "Faculty Timetable"), set())

        ay_doc = await academic_year_repo.get_by_id(tt["academicYearId"])
        st_doc = await semester_type_repo.get_by_id(tt["semesterTypeId"])
        ay_name = ay_doc.get("name", "") if ay_doc else ""
        st_name = st_doc.get("name", "") if st_doc else ""
        term_text = f"Academic Year: {ay_name} | {st_name} | Version: v{tt.get('version', 1)}"

        num_cols = len(time_slots) + 1
        subtitle = f"{fac_doc.get('name')} ({fac_doc.get('facultyCode')}, {fac_doc.get('designation')})"
        self._write_header_block(ws, inst, subtitle, term_text, num_cols)

        # Header row
        ws.row_dimensions[6].height = 28
        cell_c1 = ws.cell(row=6, column=1, value="DAY / PERIOD")
        cell_c1.font = self.font_header
        cell_c1.fill = self.fill_header
        cell_c1.alignment = self.align_center
        cell_c1.border = self.thin_border

        for col_idx, ts in enumerate(time_slots, start=2):
            s_name = ts.get("name", "")
            s_time = f"{ts.get('startTime', '')} - {ts.get('endTime', '')}"
            c = ws.cell(row=6, column=col_idx, value=f"{s_name}\n({s_time})")
            c.font = self.font_header
            c.fill = self.fill_header
            c.alignment = self.align_center
            c.border = self.thin_border

        cur_row = 7
        for wd in working_days:
            ws.row_dimensions[cur_row].height = 34
            d_cell = ws.cell(row=cur_row, column=1, value=wd.get("name", "Day"))
            d_cell.font = self.font_cell_bold
            d_cell.fill = self.fill_day
            d_cell.alignment = self.align_center
            d_cell.border = self.thin_border

            for col_idx, ts in enumerate(time_slots, start=2):
                c = ws.cell(row=cur_row, column=col_idx)
                c.border = self.thin_border
                c.alignment = self.align_center

                if ts.get("slotType") in ("BREAK", "LUNCH"):
                    c.value = ts.get("name", "")
                    c.font = self.font_meta
                    c.fill = self.fill_break
                    continue

                entry = grid_map.get((wd["id"], ts["id"]))
                if entry:
                    cid = entry.get("classId")
                    c_name = cls_map.get(cid, {}).get("name", "")
                    sid = entry.get("subjectId")
                    s_code = subj_map.get(sid, {}).get("subjectCode", entry.get("title", "")) if sid else entry.get("title", "")
                    rid = entry.get("resourceId")
                    room_str = f" [{res_map[rid].get('code')}]" if rid and rid in res_map else ""
                    c.value = f"{c_name}\n{s_code}{room_str}"
                    c.font = self.font_cell_bold
                else:
                    c.value = "—"
                    c.font = self.font_meta
            cur_row += 1

        self._apply_autofit(ws)
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    async def generate_master_workbook_excel(self, timetable_id: str) -> io.BytesIO:
        tt = await timetable_repo.get_by_id(timetable_id)
        if not tt:
            raise NotFoundException(f"Timetable with ID '{timetable_id}' not found")

        inst = await institution_settings_repo.get_settings()
        working_days = await working_day_repo.find_many({"isActive": True, "isWorkingDay": True}, sort=[("dayOrder", 1)])
        time_slots = await time_slot_repo.find_many({"isActive": True, "isTeachingSlot": True}, sort=[("slotOrder", 1)])

        classes = await class_repo.find_many(
            {"academicYearId": tt["academicYearId"], "semesterTypeId": tt["semesterTypeId"], "isActive": True},
            sort=[("name", 1)],
        )
        all_entries = await timetable_entry_repo.find_many({"timetableId": timetable_id, "isActive": True})

        subjects = await subject_repo.find_many({"isActive": True})
        subj_map = {s["id"]: s for s in subjects}
        faculty_all = await faculty_repo.find_many({"isActive": True})
        fac_map = {f["id"]: f for f in faculty_all}

        # Master lookup: (day_id, class_id, slot_id)
        master_grid = {(e.get("workingDayId"), e.get("classId"), e.get("timeSlotId")): e for e in all_entries}

        existing_sheet_names: Set[str] = set()
        wb = Workbook()

        # ----------------------------------------------------
        # SHEET 1: MASTER OVERVIEW
        # ----------------------------------------------------
        ws_master = wb.active
        ws_master.title = self._sanitize_sheet_title("Master Timetable", existing_sheet_names)

        ay_doc = await academic_year_repo.get_by_id(tt["academicYearId"])
        st_doc = await semester_type_repo.get_by_id(tt["semesterTypeId"])
        ay_name = ay_doc.get("name", "") if ay_doc else ""
        st_name = st_doc.get("name", "") if st_doc else ""
        term_text = f"Academic Year: {ay_name} | {st_name} | Master Timetable v{tt.get('version', 1)}"

        num_cols = len(time_slots) + 2
        self._write_header_block(ws_master, inst, "MASTER OVERVIEW", term_text, num_cols)

        row_idx = 6
        for wd in working_days:
            # Day separator
            ws_master.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=num_cols)
            day_banner = ws_master.cell(row=row_idx, column=1, value=f"— {wd.get('name', 'Day').upper()} —")
            day_banner.font = self.font_header
            day_banner.fill = self.fill_sub_header
            day_banner.alignment = self.align_center
            row_idx += 1

            # Header row for day
            ws_master.cell(row=row_idx, column=1, value="Class").font = self.font_header
            ws_master.cell(row=row_idx, column=1).fill = self.fill_header
            ws_master.cell(row=row_idx, column=1).alignment = self.align_center
            ws_master.cell(row=row_idx, column=1).border = self.thin_border

            for col_idx, ts in enumerate(time_slots, start=2):
                s_name = ts.get("name", "")
                s_time = f"{ts.get('startTime', '')} - {ts.get('endTime', '')}"
                c = ws_master.cell(row=row_idx, column=col_idx, value=f"{s_name}\n({s_time})")
                c.font = self.font_header
                c.fill = self.fill_header
                c.alignment = self.align_center
                c.border = self.thin_border
            row_idx += 1

            # Class rows
            for c in classes:
                cid = c["id"]
                c_name = c.get("name", "Class")
                c_cell = ws_master.cell(row=row_idx, column=1, value=c_name)
                c_cell.font = self.font_cell_bold
                c_cell.fill = self.fill_day
                c_cell.border = self.thin_border
                c_cell.alignment = self.align_center

                for col_idx, ts in enumerate(time_slots, start=2):
                    cell = ws_master.cell(row=row_idx, column=col_idx)
                    cell.border = self.thin_border
                    cell.alignment = self.align_center

                    entry = master_grid.get((wd["id"], cid, ts["id"]))
                    if entry:
                        sid = entry.get("subjectId")
                        s_code = subj_map.get(sid, {}).get("subjectCode", entry.get("title", "")) if sid else entry.get("title", "")
                        fac_codes = [fac_map[fid].get("facultyCode") or fac_map[fid].get("name", "")[:8] for fid in entry.get("facultyIds", []) if fid in fac_map]
                        fac_str = ", ".join(fac_codes)
                        cell.value = f"{s_code}\n{fac_str}" if fac_str else s_code
                        cell.font = self.font_cell_bold
                    else:
                        cell.value = "—"
                        cell.font = self.font_meta
                row_idx += 1
            row_idx += 1

        self._apply_autofit(ws_master)

        # ----------------------------------------------------
        # SHEET 2: FACULTY WORKLOAD
        # ----------------------------------------------------
        ws_wl = wb.create_sheet(title=self._sanitize_sheet_title("Faculty Workload", existing_sheet_names))
        wl_report = await report_service.get_faculty_workload_report(timetable_id=timetable_id)
        wl_headers = ["Faculty Name", "Staff Code", "Designation", "Required Hrs", "Scheduled Hrs", "Fixed Hrs", "Max Weekly Hrs", "Utilization %", "Classes Count", "Subjects Count"]
        ws_wl.append(wl_headers)
        for c_idx in range(1, len(wl_headers) + 1):
            cell = ws_wl.cell(row=1, column=c_idx)
            cell.font = self.font_header
            cell.fill = self.fill_header
            cell.border = self.thin_border
            cell.alignment = self.align_center

        for item in wl_report.items:
            ws_wl.append([
                item.facultyName,
                item.facultyCode,
                item.designation,
                item.requiredHours,
                item.scheduledHours,
                item.fixedHours,
                item.maxWeeklyHours,
                f"{item.utilizationPercent}%",
                item.classesCount,
                item.subjectsCount,
            ])
        self._apply_autofit(ws_wl)

        # ----------------------------------------------------
        # SHEET 3: SUBJECT COVERAGE
        # ----------------------------------------------------
        ws_cov = wb.create_sheet(title=self._sanitize_sheet_title("Subject Coverage", existing_sheet_names))
        cov_report = await report_service.get_subject_coverage_report(timetable_id=timetable_id)
        cov_headers = ["Class", "Subject Code", "Subject Name", "Type", "Required Hrs", "Scheduled Hrs", "Fixed Hrs", "Remaining Hrs", "Status", "Allocated Faculty"]
        ws_cov.append(cov_headers)
        for c_idx in range(1, len(cov_headers) + 1):
            cell = ws_cov.cell(row=1, column=c_idx)
            cell.font = self.font_header
            cell.fill = self.fill_header
            cell.border = self.thin_border
            cell.alignment = self.align_center

        for item in cov_report.items:
            ws_cov.append([
                item.className,
                item.subjectCode,
                item.subjectName,
                item.subjectType,
                item.requiredWeeklyHours,
                item.scheduledWeeklyHours,
                item.fixedSubjectHours,
                item.remainingHours,
                item.coverageStatus,
                ", ".join(item.facultyNames),
            ])
        self._apply_autofit(ws_cov)

        # ----------------------------------------------------
        # SHEET 4: RESOURCE UTILIZATION
        # ----------------------------------------------------
        ws_res = wb.create_sheet(title=self._sanitize_sheet_title("Resource Utilization", existing_sheet_names))
        res_report = await report_service.get_resource_utilization_report(timetable_id=timetable_id)
        res_headers = ["Resource Name", "Code", "Type", "Capacity", "Available Slots", "Occupied Slots", "Utilization %", "Classes Using"]
        ws_res.append(res_headers)
        for c_idx in range(1, len(res_headers) + 1):
            cell = ws_res.cell(row=1, column=c_idx)
            cell.font = self.font_header
            cell.fill = self.fill_header
            cell.border = self.thin_border
            cell.alignment = self.align_center

        for item in res_report.items:
            ws_res.append([
                item.resourceName,
                item.resourceCode,
                item.resourceType,
                item.capacity,
                item.availableSlots,
                item.usedSlots,
                f"{item.utilizationPercent}%",
                ", ".join(item.classesUsing),
            ])
        self._apply_autofit(ws_res)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    async def generate_faculty_workload_report_excel(
        self, academic_year_id: Optional[str] = None, semester_type_id: Optional[str] = None, timetable_id: Optional[str] = None
    ) -> io.BytesIO:
        report = await report_service.get_faculty_workload_report(
            academic_year_id=academic_year_id, semester_type_id=semester_type_id, timetable_id=timetable_id
        )
        wb = Workbook()
        ws = wb.active
        ws.title = "Faculty Workload"
        headers = ["Faculty Name", "Staff Code", "Designation", "Required Hours", "Scheduled Hours", "Fixed Hours", "Max Weekly Hours", "Utilization %", "Classes Count", "Subjects Count"]
        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=1, column=col_idx)
            c.font = self.font_header
            c.fill = self.fill_header
            c.alignment = self.align_center
            c.border = self.thin_border

        for item in report.items:
            ws.append([
                item.facultyName,
                item.facultyCode,
                item.designation,
                item.requiredHours,
                item.scheduledHours,
                item.fixedHours,
                item.maxWeeklyHours,
                f"{item.utilizationPercent}%",
                item.classesCount,
                item.subjectsCount,
            ])
        self._apply_autofit(ws)
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    async def generate_subject_coverage_report_excel(
        self, academic_year_id: Optional[str] = None, semester_type_id: Optional[str] = None, timetable_id: Optional[str] = None
    ) -> io.BytesIO:
        report = await report_service.get_subject_coverage_report(
            academic_year_id=academic_year_id, semester_type_id=semester_type_id, timetable_id=timetable_id
        )
        wb = Workbook()
        ws = wb.active
        ws.title = "Subject Coverage"
        headers = ["Class Name", "Subject Code", "Subject Name", "Type", "Required Hours", "Scheduled Hours", "Fixed Hours", "Remaining Hours", "Status", "Allocated Faculty"]
        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=1, column=col_idx)
            c.font = self.font_header
            c.fill = self.fill_header
            c.alignment = self.align_center
            c.border = self.thin_border

        for item in report.items:
            ws.append([
                item.className,
                item.subjectCode,
                item.subjectName,
                item.subjectType,
                item.requiredWeeklyHours,
                item.scheduledWeeklyHours,
                item.fixedSubjectHours,
                item.remainingHours,
                item.coverageStatus,
                ", ".join(item.facultyNames),
            ])
        self._apply_autofit(ws)
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    async def generate_resource_utilization_report_excel(
        self, academic_year_id: Optional[str] = None, semester_type_id: Optional[str] = None, timetable_id: Optional[str] = None
    ) -> io.BytesIO:
        report = await report_service.get_resource_utilization_report(
            academic_year_id=academic_year_id, semester_type_id=semester_type_id, timetable_id=timetable_id
        )
        wb = Workbook()
        ws = wb.active
        ws.title = "Resource Utilization"
        headers = ["Resource Name", "Code", "Type", "Capacity", "Available Slots", "Occupied Slots", "Utilization %", "Classes Using", "Subjects Using"]
        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=1, column=col_idx)
            c.font = self.font_header
            c.fill = self.fill_header
            c.alignment = self.align_center
            c.border = self.thin_border

        for item in report.items:
            ws.append([
                item.resourceName,
                item.resourceCode,
                item.resourceType,
                item.capacity,
                item.availableSlots,
                item.usedSlots,
                f"{item.utilizationPercent}%",
                ", ".join(item.classesUsing),
                ", ".join(item.subjectsUsing),
            ])
        self._apply_autofit(ws)
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer


excel_export_service = ExcelExportService()
